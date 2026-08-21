"""
Programmatic .docx layout for the built-in templates.

Why lay the document out in code instead of filling a .docx template
--------------------------------------------------------------------
docxtpl fills *fixed* placeholders. The second use case produces a document
whose very structure is chosen at request time - the clinician ticks which
sets and which individual fields to include - so the number of sections, their
nesting and their order all vary per request. Expressing that in a jinja2-in-
Word template means either one template per combination, or unreadable nested
`{%p for %}` tags that nobody can maintain, and neither gives control over
zebra striping, column widths or right-to-left flow.

Laying it out with python-docx instead means a new corporate look is a JSON
style profile (see templates_registry.py) rather than a hand-built binary, and
Arabic RTL is a property we set rather than a second copy of every template.
Hospitals that do need a pixel-exact letterhead can still upload a real .docx
and get the docxtpl path (renderer.py).

Everything here is pure Python: no LibreOffice, no fonts to install, so .docx
generation works identically on Linux, Windows and macOS.
"""

from __future__ import annotations

from io import BytesIO
from typing import Dict, List, Optional

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

from app.i18n import t
from app.schemas import DocField, DocRecord, DocSection, DocumentContext, FieldType
from app.templates_registry import TemplateProfile


# ---- low-level OOXML helpers --------------------------------------------


def _hex_to_rgb(value: str) -> RGBColor:
    value = (value or "000000").lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _shade(element, hex_color: str) -> None:
    """Apply a solid background fill to a table cell or paragraph."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    element.append(shd)


def _cell_shade(cell, hex_color: str) -> None:
    _shade(cell._tc.get_or_add_tcPr(), hex_color)


def _set_cell_borders(cell, *, bottom: Optional[str] = None, size: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, color in (("bottom", bottom),):
        if color is None:
            continue
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color.lstrip("#"))
        borders.append(element)


def _cell_margins(cell, top=60, bottom=60, left=100, right=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _paragraph_rtl(paragraph) -> None:
    """Mark a paragraph as right-to-left so Word orders its runs correctly."""
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    p_pr.append(bidi)


def _run_rtl(run) -> None:
    r_pr = run._r.get_or_add_rPr()
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    r_pr.append(rtl)


def _run_fonts(run, latin: str, complex_script: str) -> None:
    """
    Set both the Latin and the complex-script font.

    Word picks the *complex script* font for Arabic runs and ignores the Latin
    one, so setting only `run.font.name` leaves Arabic rendered in whatever the
    document default happens to be - which is how Arabic output ends up in a
    font with no Arabic coverage and shows as boxes.
    """
    run.font.name = latin
    r_pr = run._r.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.append(fonts)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:cs"), complex_script)


def _run_size(run, points: float) -> None:
    run.font.size = Pt(points)
    r_pr = run._r.get_or_add_rPr()
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(int(points * 2)))
    r_pr.append(sz_cs)


def _table_rtl(table) -> None:
    tbl_pr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tbl_pr.append(bidi)


def _no_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "none")
        node.set(qn("w:sz"), "0")
        borders.append(node)
    tbl_pr.append(borders)


def _add_field(paragraph, instruction: str) -> None:
    """Insert a Word field (PAGE, NUMPAGES) that recalculates on open."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# ---- the layout engine ---------------------------------------------------


class DocxLayout:
    def __init__(self, context: DocumentContext, profile: TemplateProfile) -> None:
        self.ctx = context
        self.profile = profile
        self.style = profile.resolved_style()
        self.lang = context.language.value
        self.rtl = context.direction == "rtl"
        self.font = self.style["arabic_font"] if self.rtl else self.style["font"]
        self.cs_font = self.style["arabic_font"]
        self.doc = Document()

    # -- primitives --------------------------------------------------------

    def _align(self, leading: bool = True) -> int:
        if leading:
            return WD_ALIGN_PARAGRAPH.RIGHT if self.rtl else WD_ALIGN_PARAGRAPH.LEFT
        return WD_ALIGN_PARAGRAPH.LEFT if self.rtl else WD_ALIGN_PARAGRAPH.RIGHT

    def _para(
        self,
        container=None,
        *,
        text: str = "",
        size: Optional[float] = None,
        bold: bool = False,
        italic: bool = False,
        color: Optional[str] = None,
        align: Optional[int] = None,
        space_before: float = 0,
        space_after: float = 4,
    ):
        target = container if container is not None else self.doc
        paragraph = target.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(space_before)
        fmt.space_after = Pt(space_after)
        paragraph.alignment = align if align is not None else self._align()
        if self.rtl:
            _paragraph_rtl(paragraph)
        if text:
            self._run(paragraph, text, size=size, bold=bold, italic=italic, color=color)
        return paragraph

    def _run(
        self,
        paragraph,
        text: str,
        *,
        size: Optional[float] = None,
        bold: bool = False,
        italic: bool = False,
        color: Optional[str] = None,
    ):
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        _run_fonts(run, self.font, self.cs_font)
        _run_size(run, size if size is not None else self.style["base_size"])
        if color:
            run.font.color.rgb = _hex_to_rgb(color)
        if self.rtl:
            _run_rtl(run)
        return run

    def _cell_text(self, cell, text: str, **kwargs) -> None:
        """Replace a cell's default empty paragraph rather than appending to it."""
        cell.text = ""
        paragraph = cell.paragraphs[0]
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(kwargs.pop("space_before", 2))
        fmt.space_after = Pt(kwargs.pop("space_after", 2))
        paragraph.alignment = kwargs.pop("align", None) or self._align()
        if self.rtl:
            _paragraph_rtl(paragraph)
        if text:
            self._run(paragraph, text, **kwargs)

    # -- page setup --------------------------------------------------------

    def _setup_page(self) -> None:
        section = self.doc.sections[0]
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.9)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        if self.rtl:
            sect_pr = section._sectPr
            bidi = OxmlElement("w:bidi")
            bidi.set(qn("w:val"), "1")
            sect_pr.append(bidi)

        normal = self.doc.styles["Normal"]
        normal.font.name = self.font
        normal.font.size = Pt(self.style["base_size"])
        rpr = normal.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        fonts.set(qn("w:cs"), self.cs_font)

    def _build_header(self) -> None:
        header_text = self.profile.header_text.get(self.lang) or self.ctx.facility
        if not header_text:
            return
        header = self.doc.sections[0].header
        paragraph = header.paragraphs[0]
        paragraph.alignment = self._align()
        if self.rtl:
            _paragraph_rtl(paragraph)
        self._run(
            paragraph,
            header_text,
            size=self.style["small_size"],
            bold=True,
            color=self.style["accent"],
        )

    def _build_footer(self) -> None:
        footer = self.doc.sections[0].footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if self.rtl:
            _paragraph_rtl(paragraph)

        bits: List[str] = []
        footer_text = self.profile.footer_text.get(self.lang)
        if footer_text:
            bits.append(footer_text)
        if self.ctx.options.include_generation_footer:
            bits.append(t(self.lang, "generated_footer"))
        if bits:
            self._run(
                paragraph,
                "  ·  ".join(bits),
                size=self.style["small_size"],
                color=self.style["muted"],
            )

        if self.ctx.options.include_page_numbers:
            page_para = footer.add_paragraph()
            page_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if self.rtl:
                _paragraph_rtl(page_para)
            self._run(
                page_para, f"{t(self.lang, 'page')} ", size=self.style["small_size"],
                color=self.style["muted"],
            )
            _add_field(page_para, "PAGE")
            self._run(
                page_para, f" {t(self.lang, 'of')} ", size=self.style["small_size"],
                color=self.style["muted"],
            )
            _add_field(page_para, "NUMPAGES")

    # -- document blocks ---------------------------------------------------

    def _accent_bar(self) -> None:
        if not self.style.get("show_accent_bar", True):
            return
        table = self.doc.add_table(rows=1, cols=1)
        table.autofit = True
        _no_table_borders(table)
        cell = table.rows[0].cells[0]
        _cell_shade(cell, self.style["accent"])
        _cell_margins(cell, top=20, bottom=20)
        cell.text = ""
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        self._run(cell.paragraphs[0], "", size=2)

    def _title_block(self) -> None:
        title = self.ctx.title or t(self.lang, "report_title")
        self._para(
            text=title,
            size=self.style["title_size"],
            bold=True,
            color=self.style["accent"],
            align=WD_ALIGN_PARAGRAPH.CENTER,
            space_before=10,
            space_after=2,
        )
        if self.ctx.subtitle:
            self._para(
                text=self.ctx.subtitle,
                size=self.style["base_size"] + 0.5,
                italic=True,
                color=self.style["muted"],
                align=WD_ALIGN_PARAGRAPH.CENTER,
                space_after=10,
            )

    def _patient_banner(self) -> None:
        if not self.style.get("show_patient_banner", True):
            return
        patient = self.ctx.patient
        entries: List[tuple[str, str]] = []
        if patient.display_name:
            entries.append((t(self.lang, "patient"), patient.display_name))
        if patient.identifier:
            entries.append((t(self.lang, "identifier"), patient.identifier))
        if patient.birthdate:
            entries.append((t(self.lang, "birthdate"), patient.birthdate))
        if patient.age:
            entries.append((t(self.lang, "age"), patient.age))
        if patient.gender:
            entries.append((t(self.lang, "gender"), patient.gender))
        if self.ctx.author.full_name:
            entries.append((t(self.lang, "author"), self.ctx.author.full_name))
        entries.append((t(self.lang, "generated_at"), self.ctx.generated_at))
        if self.ctx.reference:
            entries.append((t(self.lang, "reference"), self.ctx.reference))

        # Two label/value pairs per row keeps the identity block to 3-4 lines.
        rows = (len(entries) + 1) // 2
        table = self.doc.add_table(rows=rows, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _no_table_borders(table)
        if self.rtl:
            _table_rtl(table)

        for index, (label, value) in enumerate(entries):
            row, col = index // 2, (index % 2) * 2
            label_cell = table.rows[row].cells[col]
            value_cell = table.rows[row].cells[col + 1]
            for cell in (label_cell, value_cell):
                _cell_shade(cell, self.style["accent_soft"])
                _cell_margins(cell)
            self._cell_text(
                label_cell, f"{label} :", size=self.style["small_size"] + 0.5,
                bold=True, color=self.style["accent"],
            )
            self._cell_text(value_cell, value, size=self.style["small_size"] + 0.5)

        # Pad the trailing half-row so it does not render as a ragged block.
        if len(entries) % 2 == 1:
            for cell in table.rows[rows - 1].cells[2:]:
                _cell_shade(cell, self.style["accent_soft"])
                _cell_margins(cell)
                self._cell_text(cell, "")

        self._para(space_after=6)

    def _table_of_contents(self, sections: List[DocSection]) -> None:
        if not self.ctx.options.include_table_of_contents or not sections:
            return
        self._heading(t(self.lang, "table_of_contents"))
        for index, section in enumerate(sections, start=1):
            self._para(
                text=f"{index}.  {section.title}",
                size=self.style["base_size"],
                space_after=2,
            )
            for sub in section.subsections:
                self._para(
                    text=f"        •  {sub.title}",
                    size=self.style["small_size"] + 0.5,
                    color=self.style["muted"],
                    space_after=1,
                )
        self.doc.add_page_break()

    def _heading(self, text: str, level: int = 1) -> None:
        size = self.style["heading_size"] if level == 1 else self.style["heading_size"] - 1.5
        paragraph = self._para(
            text=text.upper() if level == 1 and not self.rtl else text,
            size=size,
            bold=True,
            color=self.style["accent"],
            space_before=12 if level == 1 else 8,
            space_after=3,
        )
        if level == 1:
            # A bottom rule under the heading, drawn as a paragraph border so
            # it follows the text width and survives page breaks.
            p_pr = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "8")
            bottom.set(qn("w:space"), "2")
            bottom.set(qn("w:color"), self.style["accent"].lstrip("#"))
            borders.append(bottom)
            p_pr.append(borders)

    def _visible_fields(self, fields: List[DocField]) -> List[DocField]:
        if self.ctx.options.show_empty_fields:
            return fields
        return [f for f in fields if f.value is not None and str(f.value).strip() != ""]

    def _format_value(self, field: DocField) -> str:
        value = field.value
        if value is None or str(value).strip() == "":
            return t(self.lang, "not_recorded")
        text = str(value).strip()
        if field.type == FieldType.boolean:
            key = "yes" if text.lower() in {"true", "1", "yes", "oui", "y"} else "no"
            return t(self.lang, key)
        if field.unit:
            return f"{text} {field.unit}"
        return text

    def _field_table(self, fields: List[DocField]) -> None:
        """The `Data : value` grid - two columns, zebra striped, no hard borders."""
        visible = self._visible_fields(fields)
        if not visible:
            return

        long_fields = [f for f in visible if f.type == FieldType.longtext]
        short_fields = [f for f in visible if f.type != FieldType.longtext]

        if short_fields:
            table = self.doc.add_table(rows=len(short_fields), cols=2)
            _no_table_borders(table)
            if self.rtl:
                _table_rtl(table)
            label_pct = int(self.style["label_column_percent"])
            for index, field in enumerate(short_fields):
                row = table.rows[index]
                label_cell, value_cell = row.cells[0], row.cells[1]
                label_cell.width = Cm(17 * label_pct / 100)
                value_cell.width = Cm(17 * (100 - label_pct) / 100)
                for cell in (label_cell, value_cell):
                    _cell_margins(cell)
                    if index % 2 == 1:
                        _cell_shade(cell, self.style["zebra"])
                    _set_cell_borders(cell, bottom=self.style["rule"], size=2)
                self._cell_text(
                    label_cell, f"{field.label} :", bold=True, color=self.style["muted"],
                )
                self._cell_text(
                    value_cell,
                    self._format_value(field),
                    bold=field.emphasis,
                    color=self.style["text"],
                )
            self._para(space_after=4)

        for field in long_fields:
            self._para(
                text=f"{field.label} :",
                bold=True,
                color=self.style["muted"],
                space_before=4,
                space_after=1,
            )
            for line in self._format_value(field).splitlines() or [""]:
                self._para(text=line, space_after=1)
            self._para(space_after=3)

    def _records(self, records: List[DocRecord]) -> None:
        for index, record in enumerate(records, start=1):
            heading = record.title or f"#{index}"
            paragraph = self._para(
                text=heading,
                bold=True,
                size=self.style["base_size"],
                color=self.style["text"],
                space_before=6,
                space_after=1,
            )
            _shade(paragraph._p.get_or_add_pPr(), self.style["accent_soft"])
            if record.subtitle:
                self._para(
                    text=record.subtitle,
                    italic=True,
                    size=self.style["small_size"] + 0.5,
                    color=self.style["muted"],
                    space_after=2,
                )
            self._field_table(record.fields)

    def _section(self, section: DocSection, level: int = 1) -> None:
        self._heading(section.title, level=level)
        if section.description:
            self._para(
                text=section.description,
                italic=True,
                size=self.style["small_size"] + 0.5,
                color=self.style["muted"],
                space_after=4,
            )
        self._field_table(section.fields)
        self._records(section.records)
        for sub in section.subsections:
            self._section(sub, level=min(level + 1, 3))
        if section.note:
            self._para(
                text=section.note,
                italic=True,
                size=self.style["small_size"] + 0.5,
                color=self.style["muted"],
                space_before=2,
                space_after=4,
            )

    def _image_observations(self) -> None:
        observations = self.ctx.image_observations
        if not observations:
            return
        self._heading(t(self.lang, "image_observations"))
        for observation in observations:
            label = observation.image_label or observation.study_uid or ""
            meta_bits = [
                bit
                for bit in (
                    observation.modality,
                    observation.study_date,
                    observation.author,
                    observation.observed_at,
                )
                if bit
            ]
            paragraph = self._para(
                text=label,
                bold=True,
                color=self.style["accent"],
                space_before=6,
                space_after=1,
            )
            _shade(paragraph._p.get_or_add_pPr(), self.style["accent_soft"])
            if meta_bits:
                self._para(
                    text="  ·  ".join(meta_bits),
                    italic=True,
                    size=self.style["small_size"],
                    color=self.style["muted"],
                    space_after=2,
                )
            for line in (observation.text or "").splitlines() or [""]:
                self._para(text=line, space_after=1)
            self._para(space_after=3)

    def _doctor_observation(self) -> None:
        """
        The clinician's own free text, last in the document.

        <p>Rendered verbatim: no translation, no reformatting beyond honouring the line breaks
        they typed. Placed after the imaging observations so it reads as the author's closing
        word on everything above it, which is why it is a top-level field rather than a
        Section (sections all render before the imaging block).
        """
        text = (self.ctx.doctor_observation or "").strip()
        if not text:
            return
        self._heading(t(self.lang, "doctor_observation"))
        for line in text.splitlines() or [""]:
            self._para(text=line, space_after=2)

    def _signature_block(self) -> None:
        if not self.ctx.options.include_signature_block:
            return
        self._para(space_before=16, space_after=0)
        table = self.doc.add_table(rows=2, cols=2)
        _no_table_borders(table)
        if self.rtl:
            _table_rtl(table)
        label_cell = table.rows[0].cells[1 if not self.rtl else 0]
        self._cell_text(
            label_cell,
            t(self.lang, "signature"),
            bold=True,
            size=self.style["small_size"] + 0.5,
            color=self.style["muted"],
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        line_cell = table.rows[1].cells[1 if not self.rtl else 0]
        _set_cell_borders(line_cell, bottom=self.style["muted"], size=6)
        _cell_margins(line_cell, top=340)
        self._cell_text(line_cell, "")
        if self.ctx.author.full_name:
            self._para(
                text=self.ctx.author.full_name,
                size=self.style["small_size"] + 0.5,
                align=self._align(leading=False),
                space_before=2,
            )

    def _confidentiality(self) -> None:
        if not self.ctx.options.confidentiality_notice:
            return
        self._para(space_before=10, space_after=0)
        table = self.doc.add_table(rows=1, cols=1)
        _no_table_borders(table)
        cell = table.rows[0].cells[0]
        _cell_shade(cell, self.style["accent_soft"])
        _cell_margins(cell, top=80, bottom=80)
        self._cell_text(
            cell,
            t(self.lang, "confidential"),
            size=self.style["small_size"],
            italic=True,
            color=self.style["muted"],
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    # -- entry point -------------------------------------------------------

    def build(self) -> bytes:
        self._setup_page()
        self._build_header()
        self._build_footer()
        self._accent_bar()
        self._title_block()
        self._patient_banner()

        sections = self.ctx.sections
        self._table_of_contents(sections)

        rendered_any = False
        for section in sections:
            self._section(section)
            rendered_any = True
        self._image_observations()
        self._doctor_observation()

        if (not rendered_any
                and not self.ctx.image_observations
                and not (self.ctx.doctor_observation or "").strip()):
            self._para(
                text=t(self.lang, "no_data"),
                italic=True,
                color=self.style["muted"],
                align=WD_ALIGN_PARAGRAPH.CENTER,
                space_before=20,
            )

        self._signature_block()
        self._confidentiality()

        buffer = BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()


def render_builtin_docx(context: DocumentContext, profile: TemplateProfile) -> bytes:
    return DocxLayout(context, profile).build()
